"""Tests for the two-pass sheet analysis, using real .xlsx files."""
from datetime import date, datetime
from decimal import Decimal

import pytest

from mssql_mcp_server.importer.analysis import analyze, iter_converted_rows
from mssql_mcp_server.importer.errors import ExcelImportError
from mssql_mcp_server.importer.sources import open_source


def assert_second_pass_matches(src, a):
    """Pass 2 must yield exactly the rows pass 1 counted, starting with the preview."""
    rows = list(iter_converted_rows(src, a))
    assert len(rows) == a.row_count
    assert rows[: len(a.preview_rows)] == a.preview_rows


def test_analyze_basic(make_xlsx):
    p = make_xlsx(
        {"Sales": [["id", "amount", "when"]]
         + [[i, i + 0.5, datetime(2024, 1, i)] for i in range(1, 8)]}
    )
    with open_source(p) as src:
        a = analyze(src)
        assert_second_pass_matches(src, a)
    assert a.sheet == "Sales" and a.start_cell == "A1" and a.row_count == 7
    assert a.sheets == ["Sales"]
    assert [(c.name, c.sql_type, c.nullable) for c in a.columns] == [
        ("id", "int", False),
        ("amount", "decimal(2,1)", False),
        ("when", "date", False),
    ]
    assert a.preview_rows[0] == (1, Decimal("1.5"), date(2024, 1, 1))
    assert len(a.preview_rows) == 5
    assert a.warnings == []


def test_analyze_named_sheet_and_offset_table(make_xlsx):
    p = make_xlsx(
        {
            "First": [["x"], [1]],
            "Second": [
                ["Report title"],
                [],
                [None, "id", "name"],
                [None, 1, "a"],
                [None, 2, "b"],
            ],
        }
    )
    with open_source(p) as src:
        a = analyze(src, sheet="Second")
        assert_second_pass_matches(src, a)
    assert a.sheets == ["First", "Second"]
    assert a.start_cell == "B3"
    assert a.row_count == 2
    assert [c.name for c in a.columns] == ["id", "name"]
    assert a.preview_rows == [(1, "a"), (2, "b")]
    assert a.warnings == ["1 row(s) above the table (before row 3) were skipped."]


def test_no_header_names_columns_generically(make_xlsx):
    p = make_xlsx({"S": [[1, 2], [3, 4]]})
    with open_source(p) as src:
        a = analyze(src)
        assert_second_pass_matches(src, a)
    assert [c.name for c in a.columns] == ["Column1", "Column2"]
    assert a.row_count == 2
    assert a.start_cell == "A1"


def test_ragged_rows(make_xlsx):
    # The wide row comes last, so Column3 is created late and must count the earlier
    # rows as NULL; the short row makes the second column nullable.
    p = make_xlsx({"S": [["id", "name"], [1, "a"], [3], [2, "b", "extra"]]})
    with open_source(p) as src:
        a = analyze(src)
        assert_second_pass_matches(src, a)
    assert [(c.name, c.nullable) for c in a.columns] == [
        ("id", False),
        ("name", True),
        ("Column3", True),
    ]
    assert a.row_count == 3
    assert a.preview_rows == [(1, "a", None), (3, None, None), (2, "b", "extra")]


def test_late_column_counts_earlier_rows_as_null(make_xlsx):
    p = make_xlsx({"S": [["id", "name"], [1, "a"], [2, "b", "extra"]]})
    with open_source(p) as src:
        a = analyze(src)
    assert [(c.name, c.nullable) for c in a.columns] == [
        ("id", False),
        ("name", False),
        ("Column3", True),
    ]


def test_header_only_sheet(make_xlsx):
    p = make_xlsx({"S": [["id", "name"]]})
    with open_source(p) as src:
        a = analyze(src)
        assert list(iter_converted_rows(src, a)) == []
    assert a.row_count == 0
    assert a.preview_rows == []
    assert [(c.name, c.sql_type, c.nullable) for c in a.columns] == [
        ("id", "nvarchar(255)", True),
        ("name", "nvarchar(255)", True),
    ]


def test_empty_sheet(make_xlsx):
    p = make_xlsx({"S": [[], [None, None]]})
    with open_source(p) as src:
        with pytest.raises(ExcelImportError, match="The sheet is empty"):
            analyze(src)


def test_unknown_sheet(make_xlsx):
    p = make_xlsx({"Alpha": [["a"], [1]], "Beta": [["a"], [1]]})
    with open_source(p) as src:
        with pytest.raises(ExcelImportError) as exc:
            analyze(src, sheet="Gamma")
    assert "Gamma" in str(exc.value)
    assert "Alpha" in str(exc.value) and "Beta" in str(exc.value)


def test_formula_warning(make_xlsx):
    p = make_xlsx({"S": [["a", "b"], [1, "=1+1"]]})
    with open_source(p) as src:
        a = analyze(src)
    assert len(a.warnings) == 1
    assert "1 formula cell(s)" in a.warnings[0]
    assert "no saved result" in a.warnings[0]


def test_footnote_left_of_the_table_is_not_a_data_row(make_xlsx):
    rows = (
        [["Title"], [None, "id", "name"]]
        + [[None, i, f"n{i}"] for i in range(1, 25)]
        + [["Footnote: source XYZ"]]
    )
    p = make_xlsx({"S": rows})
    with open_source(p) as src:
        a = analyze(src)
        converted = list(iter_converted_rows(src, a))
    assert a.start_cell == "B2"
    assert a.row_count == 24
    assert len(converted) == 24
    assert converted[-1] == (24, "n24")
    assert [(c.name, c.sql_type, c.nullable) for c in a.columns] == [
        ("id", "int", False),
        ("name", "nvarchar(4)", False),
    ]
    assert a.warnings == [
        "1 row(s) above the table (before row 2) were skipped.",
        "1 row(s) have content to the left of the table (before column B) "
        "that will not be imported.",
    ]


def test_content_left_of_the_table_is_warned_for_every_row(make_xlsx):
    # Beyond the sampled rows, row 26 has content on both sides: its in-table part is
    # kept and the left part dropped. Row 27 has content only on the left and is skipped.
    rows = (
        [[None, "id", "name"]]
        + [[None, i, f"n{i}"] for i in range(1, 25)]
        + [["note", 25, "n25"], ["footnote"], [None, 26, "n26"]]
    )
    p = make_xlsx({"S": rows})
    with open_source(p) as src:
        a = analyze(src)
        converted = list(iter_converted_rows(src, a))
    assert a.layout.first_col == 1
    assert a.row_count == 26
    assert converted[-2:] == [(25, "n25"), (26, "n26")]
    assert a.warnings == [
        "2 row(s) have content to the left of the table (before column B) "
        "that will not be imported."
    ]


def test_rows_above_the_table_are_warned(make_xlsx):
    rows = [["Sales report"], ["Generated 2024-02-01"], [], ["id", "name"], [1, "a"], [2, "b"]]
    p = make_xlsx({"S": rows})
    with open_source(p) as src:
        a = analyze(src)
        assert_second_pass_matches(src, a)
    assert a.start_cell == "A4"
    assert a.row_count == 2
    assert a.warnings == ["2 row(s) above the table (before row 4) were skipped."]


def test_rows_above_a_forced_header_row_are_warned(make_xlsx):
    rows = [["id", "name"], ["Real header below"], ["code", "label"], [1, "a"]]
    p = make_xlsx({"S": rows})
    with open_source(p) as src:
        a = analyze(src, header_row=3)
    assert [c.name for c in a.columns] == ["code", "label"]
    assert a.warnings == ["2 row(s) above the table (before row 3) were skipped."]


def test_forced_header_row_beyond_first_twenty_rows(make_xlsx):
    rows = [[f"note{i}"] for i in range(1, 26)] + [["id", "name"], [1, "a"]]
    p = make_xlsx({"S": rows})
    with open_source(p) as src:
        a = analyze(src, header_row=26)
        assert_second_pass_matches(src, a)
    assert a.start_cell == "A26"
    assert a.row_count == 1
    assert [c.name for c in a.columns] == ["id", "name"]


def test_forced_header_row_samples_twenty_rows_from_the_header(make_xlsx):
    # The table starts in column B at the header, but a row further down starts in
    # column A: the table's first column must reflect rows below the header too.
    rows = (
        [[f"note{i}"] for i in range(1, 25)]
        + [[None, "id", "name"], [None, 1, "a"], ["x", 2, "b"]]
    )
    p = make_xlsx({"S": rows})
    with open_source(p) as src:
        a = analyze(src, header_row=25)
        assert_second_pass_matches(src, a)
    assert a.layout.first_col == 0
    assert [c.name for c in a.columns] == ["Column1", "id", "name"]
    assert a.row_count == 2


def test_progress_called(make_xlsx):
    p = make_xlsx({"S": [["n"]] + [[i] for i in range(20001)]})
    calls = []
    with open_source(p) as src:
        a = analyze(src, progress=lambda done, total, msg: calls.append((done, total, msg)))
    assert a.row_count == 20001
    assert calls == [(10000, None, "Read 10000 rows"), (20000, None, "Read 20000 rows")]


def test_iter_converted_rows_matches(make_xlsx):
    p = make_xlsx(
        {"S": [["id", "amount"]] + [[i, i + 0.25] for i in range(1, 9)] + [[None, 1]]}
    )
    with open_source(p) as src:
        a = analyze(src)
        rows = list(iter_converted_rows(src, a))
    assert len(rows) == a.row_count == 9
    assert rows[:5] == a.preview_rows
    assert rows[0] == (1, Decimal("1.25"))
    assert rows[-1] == (None, Decimal("1"))


def test_default_sheet_skips_leading_chartsheet(tmp_path):
    import openpyxl
    from openpyxl.chart import BarChart, Reference

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    for row in [["id", "v"], [1, 2], [3, 4]]:
        ws.append(row)
    chart = BarChart()
    chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=3), titles_from_data=True)
    wb.create_chartsheet("Chart", 0).add_chart(chart)
    p = tmp_path / "chart.xlsx"
    wb.save(p)
    with open_source(p) as src:
        a = analyze(src)
    assert a.sheet == "Data"
    assert a.sheets == ["Data"]
    assert a.row_count == 2


def test_default_sheet_is_first_visible(tmp_path):
    import openpyxl

    wb = openpyxl.Workbook()
    hidden = wb.active
    hidden.title = "Lookup"
    hidden.sheet_state = "hidden"
    hidden.append(["code", "label"])
    hidden.append([1, "x"])
    shown = wb.create_sheet("Sales")
    for row in [["id", "amount"], [1, 10], [2, 20], [3, 30]]:
        shown.append(row)
    p = tmp_path / "hidden.xlsx"
    wb.save(p)
    with open_source(p) as src:
        a = analyze(src)
        assert analyze(src, sheet="Lookup").row_count == 1  # still importable by name
    assert a.sheet == "Sales"
    assert a.sheets == ["Lookup", "Sales"]  # the preview still lists every sheet
    assert a.row_count == 3


def test_default_sheet_falls_back_to_first_when_none_visible(make_xlsx, monkeypatch):
    p = make_xlsx({"A": [["x"], [1]], "B": [["y"], [2], [3]]})
    with open_source(p) as src:
        monkeypatch.setattr(src, "visible_sheet_names", lambda: [])
        a = analyze(src)
    assert a.sheet == "A"


def test_workbook_without_worksheets(make_xlsx, monkeypatch):
    p = make_xlsx({"A": [["x"], [1]]})
    with open_source(p) as src:
        monkeypatch.setattr(src, "sheet_names", lambda: [])
        with pytest.raises(ExcelImportError, match="no worksheets"):
            analyze(src)
