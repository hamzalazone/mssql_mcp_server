"""Unit tests for the Excel source layer (no database needed)."""
from datetime import datetime, time
from pathlib import Path

import openpyxl
import pytest

from mssql_mcp_server.importer.errors import ExcelImportError
from mssql_mcp_server.importer.sources import open_source

FIXTURES = Path(__file__).parent / "fixtures"


def test_xlsx_native_types(make_xlsx):
    p = make_xlsx({"S": [["a", 1, 1.5, True, datetime(2024, 2, 1)]]})
    with open_source(p) as src:
        assert src.sheet_names() == ["S"]
        row_no, cells = next(src.iter_rows("S"))
    assert row_no == 1
    assert [c.value for c in cells] == ["a", 1, 1.5, True, datetime(2024, 2, 1)]


def test_xlsx_error_cell_becomes_none(make_xlsx):
    p = make_xlsx({"S": [["#N/A", "ok"]]})
    with open_source(p) as src:
        _, cells = next(src.iter_rows("S"))
    assert [c.value for c in cells] == [None, "ok"]


def test_xlsx_time_becomes_text(make_xlsx):
    p = make_xlsx({"S": [[time(13, 45)]]})
    with open_source(p) as src:
        _, cells = next(src.iter_rows("S"))
    assert cells[0].value == "13:45:00"


def test_xlsx_number_format_exposed(make_xlsx):
    p = make_xlsx(
        {"S": [["Code"], [123]]},
        formats={("S", "A2"): "00000"},
    )
    with open_source(p) as src:
        rows = dict(src.iter_rows("S"))
    assert rows[2][0].value == 123
    assert rows[2][0].number_format == "00000"
    assert rows[1][0].number_format is None  # text cells carry no format


def test_xlsx_empty_string_and_gaps(make_xlsx):
    p = make_xlsx({"S": [["a", "", "c"], [None, None, None], ["x"]]})
    with open_source(p) as src:
        rows = {n: [c.value for c in cells] for n, cells in src.iter_rows("S")}
    assert rows[1] == ["a", None, "c"]
    assert rows[3][0] == "x"


def test_xlsx_unresolved_formula_counted(make_xlsx):
    p = make_xlsx({"S": [["=1+1", 5]]})
    with open_source(p) as src:
        assert src.unresolved_formula_count("S") == 1


def test_xlsx_no_formulas_counts_zero(make_xlsx):
    p = make_xlsx({"S": [["a", 5]]})
    with open_source(p) as src:
        assert src.unresolved_formula_count("S") == 0


def test_xls_fixture_reads_types_and_format():
    with open_source(FIXTURES / "sample.xls") as src:
        assert src.sheet_names() == ["Data"]
        rows = dict(src.iter_rows("Data"))
        assert src.unresolved_formula_count("Data") == 0
    assert [c.value for c in rows[1]] == ["Code", "Amount", "When", "Flag"]
    code, amount, when, flag = rows[2]
    assert code.value == 123.0 and isinstance(code.value, float)
    assert code.number_format == "00000"
    assert amount.value == 1.5
    assert when.value == datetime(2024, 2, 1)
    assert flag.value is True


def test_unsupported_extension(tmp_path):
    p = tmp_path / "data.csv"
    p.write_text("a,b\n1,2\n")
    with pytest.raises(ExcelImportError, match="Unsupported file type"):
        open_source(p)


def test_missing_file(tmp_path):
    with pytest.raises(ExcelImportError, match="File not found"):
        open_source(tmp_path / "nope.xlsx")


def test_directory_is_rejected(tmp_path):
    d = tmp_path / "folder.xlsx"
    d.mkdir()
    with pytest.raises(ExcelImportError, match="not a file"):
        open_source(d)


@pytest.mark.parametrize("ext", [".xlsx", ".xls"])
def test_corrupt_file(tmp_path, ext):
    p = tmp_path / f"bad{ext}"
    p.write_bytes(b"hello")
    with pytest.raises(ExcelImportError):
        open_source(p)


def test_locked_file(make_xlsx, monkeypatch):
    p = make_xlsx({"S": [["a"]]})

    def locked(*args, **kwargs):
        raise PermissionError("denied")

    monkeypatch.setattr(openpyxl, "load_workbook", locked)
    with pytest.raises(ExcelImportError, match="close it in Excel"):
        open_source(p)


def _shrink_dimension(path, ref):
    """Rewrite every worksheet's declared <dimension ref> inside the xlsx zip."""
    import re
    import zipfile

    with zipfile.ZipFile(path) as zin:
        items = [(info, zin.read(info.filename)) for info in zin.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for info, data in items:
            if info.filename.startswith("xl/worksheets/sheet"):
                data, n = re.subn(rb'<dimension ref="[^"]*"', f'<dimension ref="{ref}"'.encode(), data)
                assert n == 1
            zout.writestr(info, data)


def test_xlsx_wrong_dimension_does_not_truncate(make_xlsx):
    p = make_xlsx({"S": [[i, f"n{i}", i * 2] for i in range(1, 101)]})
    _shrink_dimension(p, "A1:B5")
    with open_source(p) as src:
        rows = list(src.iter_rows("S"))
    assert len(rows) == 100
    assert rows[-1][0] == 100
    assert [c.value for c in rows[-1][1]] == [100, "n100", 200]


def test_unresolved_formula_count_ignores_wrong_dimension(make_xlsx):
    rows = [[i, f"n{i}", i] for i in range(1, 50)] + [[50, "n50", "=1+1"]]
    p = make_xlsx({"S": rows})
    _shrink_dimension(p, "A1:B5")
    with open_source(p) as src:
        assert src.unresolved_formula_count("S") == 1


def _workbook_with_chartsheet_first(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    for row in [["id", "v"], [1, 2], [3, 4]]:
        ws.append(row)
    chart = openpyxl.chart.BarChart()
    chart.add_data(openpyxl.chart.Reference(ws, min_col=2, min_row=1, max_row=3), titles_from_data=True)
    cs = wb.create_chartsheet("Chart", 0)
    cs.add_chart(chart)
    path = tmp_path / "chart.xlsx"
    wb.save(path)
    return path


def test_chartsheets_are_not_listed(tmp_path):
    p = _workbook_with_chartsheet_first(tmp_path)
    with open_source(p) as src:
        assert src.sheet_names() == ["Data"]
        assert src.visible_sheet_names() == ["Data"]


def test_xlsx_visible_sheet_names_skip_hidden(make_xlsx, tmp_path):
    wb = openpyxl.Workbook()
    wb.active.title = "Hidden"
    wb.active.sheet_state = "hidden"
    wb.create_sheet("Shown")
    wb.create_sheet("VeryHidden").sheet_state = "veryHidden"
    p = tmp_path / "hidden.xlsx"
    wb.save(p)
    with open_source(p) as src:
        assert src.sheet_names() == ["Hidden", "Shown", "VeryHidden"]
        assert src.visible_sheet_names() == ["Shown"]


def test_xls_visible_sheet_names():
    with open_source(FIXTURES / "sample.xls") as src:
        assert src.visible_sheet_names() == ["Data"]
        src._book.sheet_by_name("Data").visibility = 1  # hidden
        assert src.visible_sheet_names() == []
