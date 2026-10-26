"""Tabular sources: read Excel workbooks and normalise cells to plain Python values.

All openpyxl/xlrd specifics stay in this module, so later layers only ever see
``RawCell`` objects holding str, int, float, bool, datetime or None.
"""
from collections.abc import Callable, Iterator
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from pathlib import Path
from typing import Any, Protocol, Self

import openpyxl
import xlrd

from .errors import ExcelImportError

SUPPORTED_EXTENSIONS = (".xlsx", ".xlsm", ".xls")

# (done, total, message); total is None when unknown.
ProgressCallback = Callable[[float, float | None, str], None]


@dataclass(frozen=True)
class RawCell:
    """One normalised cell. ``number_format`` is set for numeric cells only."""

    value: str | int | float | bool | datetime | None
    number_format: str | None = None


class TabularSource(Protocol):
    """A workbook whose sheets can be streamed row by row."""

    def sheet_names(self) -> list[str]:
        """Names of the worksheets (chartsheets excluded), in workbook order."""
        ...

    def visible_sheet_names(self) -> list[str]:
        """Names of the worksheets that are not hidden, in workbook order."""
        ...

    def iter_rows(self, sheet: str) -> Iterator[tuple[int, list[RawCell]]]:
        """Yield ``(excel_row_number_1_based, cells_from_column_A)``."""
        ...

    def unresolved_formula_count(self, sheet: str, max_rows: int = 1000) -> int:
        """Count formulas that have no cached result in the first ``max_rows`` rows."""
        ...

    def close(self) -> None: ...

    def __enter__(self) -> Self: ...

    def __exit__(self, *exc_info: object) -> None: ...


def _normalize(value: Any) -> Any:
    """Map library values onto the RawCell value types."""
    if value is None:
        return None
    if isinstance(value, str):
        return value if value != "" else None
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day)
    if isinstance(value, time):
        return value.strftime("%H:%M:%S")
    if isinstance(value, timedelta):
        return str(value)
    return value  # bool, int, float


class ExcelSource(TabularSource):
    """Reads .xlsx/.xlsm through openpyxl and .xls through xlrd."""

    def __init__(self, path: Path):
        self._path = path
        self._is_xls = path.suffix.lower() == ".xls"
        self._book: Any = None
        self._open()

    def _open(self) -> None:
        try:
            if self._is_xls:
                self._book = xlrd.open_workbook(str(self._path), formatting_info=True)
            else:
                self._book = openpyxl.load_workbook(
                    self._path, read_only=True, data_only=True
                )
        except PermissionError as exc:
            raise ExcelImportError(
                f"Cannot open '{self._path.name}': the file is locked or access is "
                "denied. If it is open in Excel, close it in Excel and try again."
            ) from exc
        except Exception as exc:  # openpyxl, xlrd and zipfile raise many types
            raise ExcelImportError(
                f"Cannot read '{self._path.name}': it is not a valid Excel workbook "
                f"({type(exc).__name__}: {exc})."
            ) from exc

    def sheet_names(self) -> list[str]:
        if self._is_xls:
            return list(self._book.sheet_names())
        return [ws.title for ws in self._book.worksheets]

    def visible_sheet_names(self) -> list[str]:
        if self._is_xls:
            return [ws.name for ws in self._book.sheets() if ws.visibility == 0]
        return [ws.title for ws in self._book.worksheets if ws.sheet_state == "visible"]

    def _xlsx_sheet(self, book: Any, sheet: str) -> Any:
        """A read-only worksheet whose declared size is ignored: a wrong <dimension>
        in the file would otherwise silently truncate rows and columns."""
        ws = book[sheet]
        ws.reset_dimensions()
        return ws

    def iter_rows(self, sheet: str) -> Iterator[tuple[int, list[RawCell]]]:
        if self._is_xls:
            yield from self._iter_xls(sheet)
        else:
            yield from self._iter_xlsx(sheet)

    def _iter_xlsx(self, sheet: str) -> Iterator[tuple[int, list[RawCell]]]:
        ws = self._xlsx_sheet(self._book, sheet)
        for row_no, row in enumerate(ws.iter_rows(min_col=1), start=1):
            yield row_no, [self._xlsx_cell(c) for c in row]

    @staticmethod
    def _xlsx_cell(cell: Any) -> RawCell:
        if getattr(cell, "data_type", None) == "e":
            return RawCell(None)
        value = _normalize(cell.value)
        fmt = None
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            fmt = cell.number_format
        return RawCell(value, fmt)

    def _iter_xls(self, sheet: str) -> Iterator[tuple[int, list[RawCell]]]:
        book = self._book
        ws = book.sheet_by_name(sheet)
        for r in range(ws.nrows):
            yield r + 1, [self._xls_cell(book, c) for c in ws.row(r)]

    @staticmethod
    def _xls_cell(book: Any, cell: Any) -> RawCell:
        ctype = cell.ctype
        if ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_ERROR):
            return RawCell(None)
        if ctype == xlrd.XL_CELL_TEXT:
            return RawCell(cell.value if cell.value != "" else None)
        if ctype == xlrd.XL_CELL_BOOLEAN:
            return RawCell(bool(cell.value))
        if ctype == xlrd.XL_CELL_DATE:
            return RawCell(xlrd.xldate_as_datetime(cell.value, book.datemode))
        fmt = book.format_map[book.xf_list[cell.xf_index].format_key].format_str
        return RawCell(float(cell.value), fmt)

    def unresolved_formula_count(self, sheet: str, max_rows: int = 1000) -> int:
        if self._is_xls:
            return 0
        formulas = openpyxl.load_workbook(self._path, read_only=True, data_only=False)
        try:
            count = 0
            cached_rows = self._xlsx_sheet(self._book, sheet).iter_rows(max_row=max_rows)
            formula_rows = self._xlsx_sheet(formulas, sheet).iter_rows(max_row=max_rows)
            for cached_row, formula_row in zip(cached_rows, formula_rows):
                for cached, formula in zip(cached_row, formula_row):
                    if (
                        cached.value is None
                        and isinstance(formula.value, str)
                        and formula.value.startswith("=")
                    ):
                        count += 1
            return count
        finally:
            formulas.close()

    def close(self) -> None:
        if self._book is None:
            return
        if self._is_xls:
            self._book.release_resources()
        else:
            self._book.close()
        self._book = None

    def __enter__(self) -> Self:
        return self

    def __exit__(self, *exc_info: object) -> None:
        self.close()


def open_source(path: Path) -> TabularSource:
    """Open ``path`` as a TabularSource, raising ExcelImportError on any problem."""
    path = Path(path)
    if path.suffix.lower() not in SUPPORTED_EXTENSIONS:
        raise ExcelImportError(
            f"Unsupported file type '{path.suffix}'. "
            f"Supported types: {', '.join(SUPPORTED_EXTENSIONS)}."
        )
    if not path.exists():
        raise ExcelImportError(f"File not found: {path}")
    if not path.is_file():
        raise ExcelImportError(f"Path is not a file: {path}")
    return ExcelSource(path)
